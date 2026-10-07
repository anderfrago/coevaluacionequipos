from datetime import timedelta
from io import BytesIO
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import load_workbook

from backend import create_app
from backend.auth import user_from_identity
from backend.models import Classroom, Evaluation, Team, User, Work, db, utcnow


@pytest.fixture()
def app():
    app = create_app(
        {
            "TESTING": True,
            "SECRET_KEY": "test-only-secret-with-at-least-32-characters",
            "SQLALCHEMY_DATABASE_URI": "sqlite://",
            "SESSION_COOKIE_SECURE": False,
            "GOOGLE_CLIENT_ID": "",
            "GOOGLE_CLIENT_SECRET": "",
        }
    )
    with app.app_context():
        db.create_all()
        db.session.add_all(
            [
                User(id=1, email="ander_frago@cuatrovientos.org", name="Admin", role="admin"),
                User(id=2, email="docente@cuatrovientos.org", name="Docente", role="teacher"),
                User(id=3, email="otro@cuatrovientos.org", name="Otro docente", role="teacher"),
                *[
                    User(
                        id=i + 10, email=f"alumno{i}@gmail.com", name=f"Alumno {i}", role="student"
                    )
                    for i in range(5)
                ],
            ]
        )
        db.session.commit()
    yield app
    with app.app_context():
        db.session.remove()
        db.drop_all()


def client_for(app, user_id):
    client = app.test_client()
    with client.session_transaction() as session:
        session["user_id"] = user_id
        session["auth_version"] = 1
        session["csrf"] = "test-csrf"
    return client


def send(client, method, path, data=None):
    return client.open(
        "/api" + path, method=method, json=data, headers={"X-CSRF-Token": "test-csrf"}
    )


def setup_team(app, size=3):
    teacher = client_for(app, 2)
    classroom = send(teacher, "POST", "/classes", {"name": "2 DAW"}).json["id"]
    response = send(
        teacher,
        "POST",
        "/works",
        {
            "classroom_id": classroom,
            "title": "Proyecto",
            "description": "Trabajo en equipo",
            "deadline": (utcnow() + timedelta(days=1)).isoformat() + "Z",
        },
    )
    assert response.status_code == 201
    work = response.json["id"]
    response = send(
        teacher,
        "POST",
        f"/works/{work}/teams",
        {
            "name": "Equipo A",
            "grade": 7,
            "emails": [f"alumno{i}@gmail.com" for i in range(size)],
        },
    )
    assert response.status_code == 201, response.json
    team = teacher.get("/api/dashboard").json["works"][0]["teams"][0]
    return teacher, work, team


def vote(app, team, evaluator, values, comment=""):
    return send(
        client_for(app, evaluator),
        "PUT",
        f"/evaluation/{team['token']}",
        {
            "allocations": {str(m["id"]): points for m, points in zip(team["members"], values)},
            "comment": comment,
        },
    )


def get_team(teacher):
    return teacher.get("/api/dashboard").json["works"][0]["teams"][0]


def test_calculation_confidentiality_and_publication(app):
    teacher, work, team = setup_team(app)
    for i in range(3):
        assert vote(app, team, 10 + i, [2, 3, 4], "Comentario confidencial").status_code == 200
    calculated = get_team(teacher)["results"]
    assert calculated["ready"]
    assert [r["grade"] for r in calculated["rows"]] == [4.67, 7.0, 9.33]
    assert calculated["rows"][0]["percentage"] == 66.67
    student = client_for(app, 10)
    before = student.get(f"/api/evaluation/{team['token']}").json["team"]
    assert "results" not in before and "evaluations" not in before and "my_result" not in before
    assert before["grade"] is None
    assert all("email" not in m for m in before["members"])
    assert (
        send(
            teacher, "POST", f"/works/{work}/state", {"action": "publish", "reviewed": True}
        ).status_code
        == 200
    )
    after = student.get(f"/api/evaluation/{team['token']}").json["team"]
    assert after["my_result"]["grade"] == 4.67
    assert "results" not in after and "evaluations" not in after
    assert vote(app, team, 10, [3, 3, 3]).status_code == 409
    assert send(teacher, "PATCH", f"/teams/{team['id']}", {"name": "Cambio"}).status_code == 409


def test_four_members_and_mr(app):
    teacher, _, team = setup_team(app, 4)
    assert team["budget"] == 12
    for i in range(4):
        response = vote(app, team, 10 + i, [3, 3, 3, 3])
        assert response.status_code == 200
        assert "mismos puntos" in response.json["warnings"][0]
    assert all(
        row["grade"] == 7 and row["percentage"] == 100
        for row in get_team(teacher)["results"]["rows"]
    )


def test_mh_and_zero_scores(app):
    teacher, _, team = setup_team(app)
    for i in range(3):
        response = vote(app, team, 10 + i, [9, 0, 0])
        assert len(response.json["warnings"]) == 2
    rows = get_team(teacher)["results"]["rows"]
    assert rows[0]["grade"] == 10 and rows[0]["raw_grade"] == 21 and rows[0]["mh"]
    assert rows[1]["grade"] == 0 and not rows[1]["mh"]


@pytest.mark.parametrize(
    "values", [[2.5, 2.5, 4], [-1, 5, 5], [True, 4, 4], [3, 3, 2], [10, 0, -1]]
)
def test_invalid_allocations(app, values):
    _, _, team = setup_team(app)
    assert vote(app, team, 10, values).status_code == 400


def test_pending_results_and_replacement(app):
    teacher, work, team = setup_team(app)
    assert vote(app, team, 10, [3, 3, 3]).status_code == 200
    assert vote(app, team, 10, [2, 3, 4]).status_code == 200
    data = get_team(teacher)
    assert data["submitted"] == 1 and not data["results"]["ready"]
    assert all(row["grade"] is None for row in data["results"]["rows"])
    assert send(teacher, "POST", f"/works/{work}/state", {"action": "publish"}).status_code == 409
    student = client_for(app, 10)
    response = send(
        student, "PUT", f"/evaluation/{team['token']}", {"allocations": {"10": 9}, "comment": ""}
    )
    assert response.status_code == 400


def test_access_ownership_and_csrf(app):
    _, work, team = setup_team(app)
    assert app.test_client().get("/api/dashboard").status_code == 401
    student = client_for(app, 10)
    assert student.post("/api/classes", json={"name": "Forbidden"}).status_code == 403
    assert send(student, "POST", "/classes", {"name": "Forbidden"}).status_code == 403
    assert student.get("/api/users").status_code == 403
    assert student.get(f"/api/works/{work}/export").status_code == 403
    assert client_for(app, 14).get(f"/api/evaluation/{team['token']}").status_code == 404
    assert client_for(app, 3).get(f"/api/works/{work}/export").status_code == 403
    assert send(client_for(app, 3), "DELETE", f"/teams/{team['id']}").status_code == 403
    assert client_for(app, 3).get("/api/dashboard").json["works"] == []


def test_expiry_and_roster_lock(app):
    teacher, work, team = setup_team(app)
    vote(app, team, 10, [2, 3, 4])
    response = send(
        teacher,
        "PATCH",
        f"/teams/{team['id']}",
        {
            "name": "A",
            "grade": 7,
            "emails": ["alumno0@gmail.com", "alumno1@gmail.com", "alumno4@gmail.com"],
        },
    )
    assert response.status_code == 409
    with app.app_context():
        db.session.get(Work, work).deadline = utcnow() - timedelta(seconds=1)
        db.session.commit()
    assert vote(app, team, 11, [2, 3, 4]).status_code == 409


def test_duplicate_membership(app):
    teacher, work, _ = setup_team(app)
    response = send(
        teacher,
        "POST",
        f"/works/{work}/teams",
        {"name": "B", "grade": None, "emails": ["alumno0@gmail.com", "alumno4@gmail.com"]},
    )
    assert response.status_code == 409
    assert len(teacher.get("/api/dashboard").json["works"][0]["teams"]) == 1


def test_export_real_xlsx_pending_and_formula_injection(app):
    teacher, work, team = setup_team(app)
    vote(app, team, 10, [2, 3, 4], '=HYPERLINK("https://example.com")')
    response = teacher.get(f"/api/works/{work}/export")
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.data))
    assert book.sheetnames == ["Resultados", "Repartos confidenciales", "Criterios"]
    assert book["Resultados"]["I2"].value is None
    assert book["Resultados"]["L2"].value == "Pendiente"
    assert book["Repartos confidenciales"]["G2"].data_type == "s"
    assert book["Repartos confidenciales"]["G2"].value.startswith("'=")
    assert book["Resultados"].freeze_panes == "A2"


def test_user_crud_and_revocation(app):
    admin = client_for(app, 1)
    data = {"name": "Nueva", "email": "nueva@gmail.com", "role": "student", "active": True}
    created = send(admin, "POST", "/users", data)
    assert created.status_code == 201
    user_id = created.json["id"]
    assert send(admin, "PATCH", f"/users/{user_id}", {**data, "role": "admin"}).status_code == 400
    active_session = client_for(app, user_id)
    assert send(admin, "PATCH", f"/users/{user_id}", {**data, "active": False}).status_code == 200
    assert active_session.get("/api/dashboard").status_code == 401
    assert send(admin, "PATCH", f"/users/{user_id}", {**data, "active": True}).status_code == 200
    assert active_session.get("/api/dashboard").status_code == 401
    assert send(admin, "DELETE", f"/users/{user_id}").status_code == 200
    assert send(admin, "DELETE", "/users/1").status_code == 400


def test_existing_history_preserved_and_domain_restriction(app):
    setup_team(app)
    admin = client_for(app, 1)
    assert send(admin, "DELETE", "/users/10").status_code == 409
    response = send(
        admin,
        "POST",
        "/users",
        {"name": "Fake", "email": "x@cuatrovientos.org.evil.test", "role": "teacher"},
    )
    assert response.status_code == 400


def test_oauth_verified_identity_and_no_role_escalation(app):
    with app.app_context():
        db.session.add(User(email="registro@gmail.com", name="Registro", role="student"))
        db.session.commit()
        user = user_from_identity(
            {
                "email": "registro@gmail.com",
                "email_verified": True,
                "sub": "google-123",
                "name": "Registro",
                "role": "admin",
            }
        )
        assert user.role == "student"
        first_id = user.id
        assert (
            user_from_identity(
                {"email": "registro@gmail.com", "email_verified": True, "sub": "google-123"}
            ).id
            == first_id
        )
        with pytest.raises(Exception) as error:
            user_from_identity({"email": "admin@gmail.com", "email_verified": False, "sub": "bad"})
        assert error.value.code == 403
        admin = user_from_identity(
            {
                "email": "ander_frago@cuatrovientos.org",
                "email_verified": True,
                "sub": "admin-google",
                "hd": "cuatrovientos.org",
            }
        )
        assert admin.role == "admin"
    response = app.test_client().get("/auth/google")
    assert response.status_code == 302 and "configuration" in response.location


def test_mail_configuration_and_individual_recipients(app):
    teacher, _, team = setup_team(app)
    app.config.update(SMTP_USERNAME="", SMTP_PASSWORD="", MAIL_FROM="")
    assert send(teacher, "POST", f"/teams/{team['id']}/invite", {}).status_code == 503
    app.config.update(
        SMTP_USERNAME="sender@gmail.com", SMTP_PASSWORD="test-only", MAIL_FROM="sender@gmail.com"
    )
    smtp = MagicMock()
    with patch("backend.mail.smtplib.SMTP") as connection:
        connection.return_value.__enter__.return_value = smtp
        response = send(teacher, "POST", f"/teams/{team['id']}/invite", {})
    assert len(response.json["sent"]) == 3
    assert smtp.send_message.call_count == 3
    for call in smtp.send_message.call_args_list:
        message = call.args[0]
        assert message["To"].count("@") == 1
        assert message["Cc"] is None
        assert team["token"] in message.get_body(preferencelist=("plain",)).get_content()


def test_delete_work_cascades(app):
    teacher, work, team = setup_team(app)
    vote(app, team, 10, [2, 3, 4])
    assert send(teacher, "DELETE", f"/works/{work}").status_code == 200
    with app.app_context():
        assert db.session.query(Evaluation).count() == 0
        assert db.session.query(Team).count() == 0
        assert db.session.query(Classroom).count() == 1
