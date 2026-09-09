from io import BytesIO

from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .grading import results, warnings


def safe(value):
    # Evita interpretar datos aportados por usuarios como fórmulas de Excel.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def export_work(work):
    book = Workbook()
    summary = book.active
    summary.title = "Resultados"
    summary.append(
        [
            "Trabajo",
            "Equipo",
            "Alumno",
            "Correo",
            "Nota grupal",
            "Puntos X1",
            "% nota",
            "Nota sin tope",
            "Nota individual",
            "MH",
            "MR en el equipo",
            "Estado",
        ]
    )
    detail = book.create_sheet("Repartos confidenciales")
    detail.append(
        [
            "Equipo",
            "Evaluador",
            "Destinatario",
            "Puntos",
            "MR",
            "Motivos MR",
            "Comentario",
            "Actualizado (UTC)",
        ]
    )
    for team in work.teams:
        calculated = results(team)
        mr = any(warnings(e.allocations) for e in team.evaluations)
        for row in calculated["rows"]:
            summary.append(
                [
                    safe(v)
                    for v in [
                        work.title,
                        team.name,
                        row["name"],
                        row["email"],
                        float(team.grade) if team.grade is not None else None,
                        row["points"],
                        row["percentage"] / 100 if row["percentage"] is not None else None,
                        row["raw_grade"],
                        row["grade"],
                        "MH" if row["mh"] else "",
                        "MR" if mr else "",
                        "Publicado"
                        if work.published
                        else "Sin publicar"
                        if calculated["ready"]
                        else "Pendiente",
                    ]
                ]
            )
        names = {str(m.user_id): m.user.email for m in team.members}
        for evaluation in team.evaluations:
            reasons = warnings(evaluation.allocations)
            for recipient, points in evaluation.allocations.items():
                detail.append(
                    [
                        safe(v)
                        for v in [
                            team.name,
                            names[str(evaluation.evaluator_id)],
                            names[recipient],
                            points,
                            "MR" if reasons else "",
                            " ".join(reasons),
                            evaluation.comment,
                            evaluation.updated_at.isoformat(),
                        ]
                    ]
                )
    notes = book.create_sheet("Criterios")
    for row in [
        ["Concepto", "Regla"],
        ["Puntos por evaluador", "3 × número de integrantes. Incluye autoevaluación."],
        ["Factor", "X1 / (3 × número de integrantes). Se calcula con precisión decimal."],
        ["Nota", "Mínimo entre 10 y nota grupal × factor; redondeo final a 2 decimales."],
        ["MR", "Ceros, todos los puntos a una persona o reparto igualitario. No resta nota."],
        ["MH", "Nota sin limitar superior a 10. No se activa por redondear a 10."],
        ["Pendiente", "Faltan evaluaciones o nota grupal. Las notas permanecen vacías."],
        [
            "Confidencial",
            "Uso exclusivo del profesorado autorizado. Contiene comentarios y repartos.",
        ],
    ]:
        notes.append(row)
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 30
        for cell in sheet[1]:
            cell.font = Font(name="Segoe UI", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F4285")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Segoe UI", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F0F5FC")
        for i, column in enumerate(sheet.columns, 1):
            width = min(65, max(15, max(len(str(c.value or "")) for c in column) + 2))
            sheet.column_dimensions[get_column_letter(i)].width = width
    for row in summary.iter_rows(min_row=2):
        for index in (4, 7, 8):
            row[index].number_format = "0.00"
        row[6].number_format = "0.00%"
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name=f"coevaluacion-trabajo-{work.id}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
